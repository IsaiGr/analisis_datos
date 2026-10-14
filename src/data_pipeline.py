from __future__ import annotations

from io import BytesIO
from pathlib import Path
from dataclasses import dataclass
from numbers import Number
import numpy as np
import pandas as pd
from openpyxl import load_workbook

FORECAST_FILE_HINT = "inventario.xlsx"
TRANSIT_FILE_HINT = "transitos.xlsx"
DEMAND_FILE_HINT = "ordenes.xlsx"
ISSUE_COLUMNS = ["source", "sheet", "row", "column", "raw", "issue"]
COVERAGE_PRIORITY_MONTHS = 4.0
COVERAGE_PRIORITY_LABEL = f"Prioritaria: menos de {COVERAGE_PRIORITY_MONTHS:g} meses"


@dataclass(frozen=True)
class IssueLog:
    """Immutable provenance, shared safely across pandas copies and selections."""
    entries: tuple

    @classmethod
    def from_records(cls, records):
        return cls(tuple(tuple(record.get(key) for key in ISSUE_COLUMNS) for record in records))

    def __iter__(self):
        return (dict(zip(ISSUE_COLUMNS, entry)) for entry in self.entries)

    def __deepcopy__(self, memo):
        return self


class MissingSheetError(ValueError):
    pass


def _source(source):
    if isinstance(source, (str, Path)):
        return str(source)
    return BytesIO(source) if isinstance(source, bytes) else source


def file_signature(path):
    """Invalidate cached results when a local workbook is saved or replaced."""
    stat = Path(path).stat()
    return stat.st_mtime_ns, stat.st_size


def clean_text(values):
    text = values.astype("string").str.strip().replace({"": pd.NA, "nan": pd.NA, "None": pd.NA})
    return text.mask(text.isin(["#N/A", "#VALUE!", "#REF!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"]))


def clean_item(values):
    # Excel numeric 123.0 and text '123' are the same ID; text '00123' stays intact.
    return clean_text(values.map(lambda value: str(int(value)) if isinstance(value, Number) and not pd.isna(value) and np.isfinite(value) and float(value).is_integer() else value))


def _read_cached_sheet(source, sheet_name, header_row=0, max_col=None):
    """Preserve Excel row numbers and errors before numeric conversion."""
    wb = load_workbook(_source(source), read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise MissingSheetError(f"No existe la hoja '{sheet_name}'. Hojas disponibles: {wb.sheetnames}")
        all_rows = list(wb[sheet_name].iter_rows(values_only=True, max_col=max_col))
    finally:
        wb.close()
    if len(all_rows) <= header_row:
        raise ValueError(f"La hoja '{sheet_name}' no contiene encabezados y datos.")
    headers = [str(x).strip() if x is not None else f"col_{i}" for i, x in enumerate(all_rows[header_row])]
    if len(headers) != len(set(headers)):
        raise ValueError(f"La hoja '{sheet_name}' contiene encabezados repetidos; revisa la estructura.")
    df = pd.DataFrame(all_rows[header_row + 1:], columns=headers)
    populated = df.notna().any(axis=1)
    df["source_row"] = np.arange(header_row + 2, header_row + 2 + len(df))
    df["source_sheet"] = sheet_name
    df = df.loc[populated].reset_index(drop=True)
    issues = []
    for col in headers:
        for index in df.index[df[col].astype("string").str.startswith("#", na=False)]:
            issues.append(dict(source="", sheet=sheet_name, row=int(df.at[index, "source_row"]),
                               column=col, raw=str(df.at[index, col]), issue="Error de Excel"))
    df.attrs["issues"] = IssueLog.from_records(issues)
    return df


def _normalize(df, rename, numeric, source_name):
    out = df.rename(columns=rename).copy()
    issues = [dict(issue, source=source_name) for issue in df.attrs.get("issues", [])]
    reverse = {value: key for key, value in rename.items()}
    sheet = out["source_sheet"].iloc[0] if len(out) else "Sin registros"
    for col in dict.fromkeys(rename.values()):
        if col not in out:
            out[col] = np.nan if col in numeric else pd.NA
            issues.append(dict(source=source_name, sheet=sheet, row=None,
                               column=reverse[col], raw="", issue="Columna ausente"))
            continue
        if col not in numeric:
            continue
        raw = out[col]
        text = raw.astype("string").str.strip()
        missing = raw.isna() | text.isin(["", "-", "—", "N/A", "n/a"])
        numbers = pd.to_numeric(raw.mask(missing), errors="coerce").astype(float)
        invalid = ~missing & (numbers.isna() | ~np.isfinite(numbers))
        negative = numbers.lt(0)
        for mask, reason in [(missing, "Sin dato numérico"), (invalid, "Número inválido"), (negative, "Valor negativo por revisar")]:
            for index in out.index[mask]:
                issues.append(dict(source=source_name, sheet=sheet, row=int(out.at[index, "source_row"]),
                                   column=reverse[col], raw="" if pd.isna(raw.at[index]) else str(raw.at[index]), issue=reason))
        out[col] = numbers.mask(invalid | negative)
    for col in ["item", "family", "category", "origin", "priority", "abc_class", "status", "invoice", "description"]:
        if col in out:
            out[col] = clean_item(out[col]) if col == "item" else clean_text(out[col])
    out["source"] = source_name
    out.attrs["issues"] = IssueLog.from_records(issues)
    return out


def find_file(folder, hint):
    exact = folder / hint
    if exact.exists():
        return exact
    key = hint.split(" ")[0].lower()
    return next((p for p in folder.glob("*.xlsx") if key in p.name.lower()), None)


def load_forecast(source):
    df = _read_cached_sheet(source, "1. FORECAST", header_row=3, max_col=72)
    rename = {
        "Item": "item", "Descripción": "description", "Familia": "family",
        "Categoría": "category", "Modelo del Producto": "customer_model",
        "Origen": "origin", "Key\nComponent": "key_component",
        "Precio Compra (MXN)": "unit_cost_mxn", "Inventario\nActual": "inventory",
        "Compras\nen Camino": "on_order", "Prom.\nMensual (3M)": "monthly_demand",
        "Lead\nTime": "lead_time", "Punto de\nReorden (pza)": "reorder_point",
        "¿Reordenar?": "reorder_flag", "Cantidad Final de Compra": "recommended_qty",
        "Monto Total de Compra": "purchase_amount", "Semáforo\nPrioridad": "priority",
        "Salud\nInventario": "inventory_health", "Clase\nABC": "abc_class",
        "Estatus de Soporte": "support_status", "Estatus de Certificado NOM": "nom_status",
        "Decisión Final de Compra": "purchase_decision", "Valor de Inventario": "inventory_value",
        "Avg Back\nOrder (3M)": "avg_back_order_3m",
    }
    numeric = {"unit_cost_mxn", "inventory", "on_order", "monthly_demand", "lead_time", "reorder_point",
               "recommended_qty", "purchase_amount", "inventory_value", "avg_back_order_3m"}
    return _normalize(df, rename, numeric, "forecast")


def load_demand(source):
    df = _read_cached_sheet(source, "Free Sales", header_row=1, max_col=20)
    rename = {
        "Week": "week", "Posting Date": "posting_date", "Created Date": "created_date",
        "Sales Order": "sales_order", "User Status": "status", "Quantity": "quantity",
        "Parts ID": "item", "Description": "description", "Customer Model": "customer_model",
        "Factory Model": "factory_model", "Key Component": "key_component",
    }
    df = _normalize(df, rename, {"quantity"}, "demand")
    issues = list(df.attrs["issues"])
    for col in ["posting_date", "created_date"]:
        raw = df[col]
        safe = raw.mask(raw.map(lambda value: isinstance(value, (int, float))))
        df[col] = pd.to_datetime(safe, errors="coerce", format="mixed")
        for index in df.index[df[col].isna()]:
            issues.append(dict(source="demand", sheet="Free Sales", row=int(df.at[index, "source_row"]),
                               column=col, raw=str(raw.at[index]), issue="Fecha ausente o inválida"))
    df.attrs["issues"] = IssueLog.from_records(issues)
    return df


def load_transits(source):
    try:
        df = _read_cached_sheet(source, "INBOUND", header_row=0, max_col=16)
    except MissingSheetError:
        # Only a missing sheet warrants fallback, not a malformed INBOUND column.
        if hasattr(source, "seek"):
            source.seek(0)
        df = _read_cached_sheet(source, "DASHBOARD", header_row=4, max_col=16)
        df["status"] = clean_text(df.iloc[:, 0])
        valid = {"Completed", "In Transit", "Pending Dispatch", "Partial Inbound", "Collect", "Not received", "Partial Delivered", "Pending Inbound", "Validate"}
        df = df.loc[df["status"].isin(valid)].copy()
        rename = {"Cantidad de facturas": "invoice_count", "Monto de factura": "invoice_amount_usd", "Cantidad de piezas": "planned_units", "Piezas recibidas": "received_units", "Facturas pendientes": "pending_invoices", "Piezas no recibidas": "unreceived_units", "Total pendientes": "total_pending"}
        df = _normalize(df, rename, set(rename.values()), "transits")
        return df[["status", *rename.values(), "source", "source_row", "source_sheet"]]
    rename = {
        "PART. NO.": "item", "DESCRIPTION": "description", "QTY": "quantity",
        "UNIT PRICE (USD)": "unit_price_usd", "CONTAINER NO.  / GUIDE": "container",
        "FACTORY": "origin", "PRODUCT LINE": "category", "INVOICE NO.": "invoice",
        "ETA": "eta", "STATUS": "status",
    }
    df = _normalize(df, rename, {"quantity", "unit_price_usd"}, "transits")
    df["invoice_amount_usd"] = df["quantity"] * df["unit_price_usd"]
    df["status"] = df["status"].str.upper().replace({
        "ENTERED INTO SYSTEM": "Pending Inbound", "PENDING INBOUND": "Pending Inbound",
        "IN TRANSIT": "In Transit", "COMPLETED": "Completed", "PARTIAL INBOUND": "Partial Inbound",
    })
    return df


def catalog_matches(items, forecast):
    counts = clean_item(forecast["item"]).value_counts()
    keys = clean_item(items)
    size = keys.map(counts).fillna(0)
    return pd.Series(np.select([keys.isna(), size.eq(0), size.gt(1)],
                              ["Sin código", "Sin coincidencia", "Código repetido"], default="Relación única"), index=items.index)


def add_forecast_context(frame, forecast):
    """Only unique codes supply catalog attributes. Never choose a duplicate."""
    if "item" not in frame:
        return frame.copy()
    out = frame.copy()
    out["item"] = clean_item(out["item"])
    out["catalog_match"] = catalog_matches(out["item"], forecast)
    attrs = [c for c in ["family", "category", "origin", "priority", "abc_class", "key_component"] if c in forecast]
    lookup = forecast[["item"] + attrs].copy()
    lookup["item"] = clean_item(lookup["item"])
    lookup = lookup.loc[lookup["item"].notna() & ~lookup["item"].duplicated(keep=False)].set_index("item")
    for col in attrs:
        catalog_value = out["item"].map(lookup[col])
        out[col] = clean_text(out[col]).fillna(catalog_value) if col in out else catalog_value
    return out


def strict_sum(values):
    """An incomplete total is unknown. An empty selection or measured zero is zero."""
    if values.empty:
        return 0.0
    numbers = pd.to_numeric(values, errors="coerce")
    return float(numbers.sum()) if numbers.notna().all() and np.isfinite(numbers).all() else float("nan")


def pending_demand(frame):
    if not {"status", "quantity"}.issubset(frame):
        return {"quantity": np.nan, "rows": 0, "unknown_status": len(frame)}
    status = clean_text(frame["status"]).str.casefold().str.replace(r"[\s_-]+", "", regex=True)
    pending = frame.loc[status.eq("backorder").fillna(False)]
    total = strict_sum(pending["quantity"]) if frame.empty or status.notna().any() else np.nan
    return {"quantity": total, "rows": len(pending), "unknown_status": int(status.isna().sum())}


def apply_filters(frame, selections, options, risk_items=None, coverage_items=None):
    out = frame.copy()
    for col, selected in selections.items():
        if not selected:
            return out.iloc[:0]
        if set(selected) != set(options[col]):
            if col not in out:
                return out.iloc[:0]
            out = out.loc[out[col].fillna("Sin relación / sin dato").isin(selected)]
    for selected_items in [risk_items, coverage_items]:
        if selected_items is not None:
            if "item" not in out:
                return out.iloc[:0]
            out = out.loc[out["item"].isin(selected_items)]
    return out


def add_coverage_indicators(frame):
    """Physical stock / adjusted monthly demand, independent of delivery time.

    Unknown or zero demand is not zero months. Duplicate codes remain excluded
    even when a later filter leaves only one of their original rows.
    """
    out = frame.copy()
    items = clean_item(out["item"])
    eligible = items.notna() & ~items.duplicated(keep=False)
    if "catalog_eligible" in out:
        eligible &= out["catalog_eligible"].fillna(False)
    inventory = pd.to_numeric(out["inventory"], errors="coerce")
    demand = pd.to_numeric(out["monthly_demand"], errors="coerce")
    known = (np.isfinite(inventory) & np.isfinite(demand) & inventory.ge(0) & demand.ge(0)).fillna(False)
    coverage = inventory / demand.where(demand.gt(0))
    valid = (eligible & known & demand.gt(0) & np.isfinite(coverage)).fillna(False)
    out["coverage_calculable"] = valid
    out["coverage_months_calc"] = coverage.where(valid)
    out["coverage_priority"] = pd.Series(pd.NA, index=out.index, dtype="boolean")
    out.loc[valid, "coverage_priority"] = coverage.loc[valid].lt(COVERAGE_PRIORITY_MONTHS)
    out["coverage_status"] = np.select(
        [~eligible, (eligible & known & demand.eq(0)).fillna(False).to_numpy(dtype=bool), ~valid.to_numpy(dtype=bool), out["coverage_priority"].fillna(False).to_numpy(dtype=bool)],
        ["Código por revisar", "Sin consumo registrado", "Sin datos para calcular", COVERAGE_PRIORITY_LABEL],
        default=f"Cobertura de {COVERAGE_PRIORITY_MONTHS:g} meses o más",
    )
    return out


def sort_by_coverage_priority(frame):
    """Give the business coverage rule precedence over the Excel purchase amount."""
    out = frame.copy()
    out["_coverage_rank"] = out["coverage_priority"].map({True: 0, False: 1}).fillna(2)
    keys = ["_coverage_rank", "coverage_months_calc"]
    ascending = [True, True]
    for column, direction in [("shortage_gap", False), ("purchase_amount", False), ("item", True)]:
        if column in out:
            keys.append(column)
            ascending.append(direction)
    return out.sort_values(keys, ascending=ascending, na_position="last", kind="stable").drop(columns="_coverage_rank")


def coverage_text(value):
    """Two decimals without rounding a priority up to the four-month boundary."""
    if pd.isna(value) or not np.isfinite(value) or value < 0:
        return "No calculable"
    if 0 < value < 0.005:
        return "< 0.01 meses"
    if value < COVERAGE_PRIORITY_MONTHS <= round(value, 2):
        return f"< {COVERAGE_PRIORITY_MONTHS:g} meses"
    return f"{value:,.2f}".rstrip("0").rstrip(".") + " meses"


def enrich_forecast(df):
    out = df.copy()
    out["item"] = clean_item(out["item"])
    out["duplicate_item"] = out["item"].notna() & out["item"].duplicated(keep=False)
    out["catalog_eligible"] = out["item"].notna() & ~out["duplicate_item"]
    fields = ["inventory", "monthly_demand", "lead_time"]
    for col in fields + ["recommended_qty", "unit_cost_mxn", "purchase_amount"]:
        if col not in out:
            out[col] = np.nan
    inputs = out[fields].apply(pd.to_numeric, errors="coerce")
    valid_inputs = (inputs.notna() & np.isfinite(inputs) & inputs.ge(0)).all(axis=1)
    out["risk_calculable"] = out["catalog_eligible"] & valid_inputs
    # Source AZ is MONTHS. Gross exposure is not a net purchase proposal.
    out["demand_during_lead_time"] = (inputs["monthly_demand"] * inputs["lead_time"]).where(out["risk_calculable"])
    out["shortage_gap"] = (out["demand_during_lead_time"] - inputs["inventory"]).clip(lower=0)
    out = add_coverage_indicators(out)
    out["risk_note"] = np.select(
        [out["item"].isna(), out["duplicate_item"], ~valid_inputs],
        ["Sin código: no calculable", "Código repetido: confirmar registro correcto", "Faltan datos válidos: no calculable"],
        default="Estimación bruta; confirmar pedidos en camino, compromisos y fechas",
    )
    expected = out["recommended_qty"] * out["unit_cost_mxn"]
    comparable = expected.notna() & out["purchase_amount"].notna()
    out["purchase_amount_consistent"] = pd.Series(pd.NA, index=out.index, dtype="boolean")
    out.loc[comparable, "purchase_amount_consistent"] = np.isclose(expected[comparable], out.loc[comparable, "purchase_amount"], rtol=0, atol=0.02)
    out["purchase_review"] = np.select(
        [~out["catalog_eligible"], ~comparable, out["purchase_amount_consistent"].eq(False).fillna(False).to_numpy(dtype=bool)],
        ["Revisar código antes de comprar", "Sin datos para comprobar importe", "Revisar: importe distinto a cantidad × costo"],
        default="Importe consistente; propuesta del Excel pendiente de validación operativa",
    )
    return out


def data_issues(*frames):
    return pd.DataFrame([issue for frame in frames for issue in frame.attrs.get("issues", [])], columns=ISSUE_COLUMNS)


def validate_inputs(forecast, demand, transits):
    checks = []
    def add(name, condition, detail):
        checks.append({"check": name, "status": "OK" if condition else "REVISAR", "detail": detail})
    for label, frame in [("Forecast", forecast), ("Demanda", demand), ("Tránsitos", transits)]:
        add(f"{label} cargado", not frame.empty, f"{len(frame):,} registros")
    duplicate = forecast.loc[forecast["duplicate_item"]]
    add("Códigos repetidos", duplicate.empty, f"{duplicate['item'].nunique():,} códigos / {len(duplicate):,} filas excluidas de indicadores de inventario")
    no_code = sum(int(frame["item"].isna().sum()) for frame in (forecast, demand, transits) if "item" in frame)
    add("Registros sin código", no_code == 0, f"{no_code:,} registros; no se relacionan con el catálogo")
    unknown = int((~forecast["risk_calculable"]).sum())
    add("Datos para estimar faltantes", unknown == 0, f"{unknown:,} filas no calculables (incluye códigos repetidos)")
    inconsistent = int(forecast["purchase_amount_consistent"].eq(False).sum())
    add("Importe de compra = cantidad × costo", inconsistent == 0, f"{inconsistent:,} diferencias; no se corrigen automáticamente")
    for label, frame in [("Demanda", demand), ("Tránsitos", transits)]:
        unmatched = int(frame["catalog_match"].ne("Relación única").sum()) if "catalog_match" in frame else len(frame)
        add(f"Relación única de {label.lower()}", unmatched == 0, f"{unmatched:,} registros sin relación única")
    issues = data_issues(forecast, demand, transits)
    for kind in ["Columna ausente", "Error de Excel", "Número inválido", "Valor negativo por revisar", "Sin dato numérico", "Fecha ausente o inválida"]:
        count = int(issues["issue"].eq(kind).sum())
        add(kind, count == 0, f"{count:,} incidencias; consulta hoja, fila y valor original abajo")
    return {"checks": checks, "error_count": int(issues["issue"].eq("Error de Excel").sum()), "issues": issues}
